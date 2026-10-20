#!/usr/bin/env python3
"""PPTX ingredients for reconstruction, derived from Frontend Slides (MIT).

Retains group hierarchy, positions, styled text runs, tables, chart series,
images and notes. Does not render or claim fidelity for unsupported objects.
"""
import argparse
import json
from pathlib import Path
import sys


def extract(source, output):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    output=Path(output).resolve();assets=output/'assets';assets.mkdir(parents=True,exist_ok=True)
    prs=Presentation(source)
    warnings=[]
    def shape_data(shape, trail):
        row={'id':trail,'name':shape.name,'shape_type':int(shape.shape_type),
             'box_emu':{'x':shape.left,'y':shape.top,'width':shape.width,'height':shape.height},
             'rotation':shape.rotation}
        if shape.shape_type==MSO_SHAPE_TYPE.GROUP:
            row['children']=[shape_data(child,f'{trail}-{i+1}') for i,child in enumerate(shape.shapes)]
            row['coordinate_note']='Group child coordinates retain source coordinate system; apply group transform for flattened placement.'
        if shape.has_text_frame:
            row['text']=shape.text
            row['paragraphs']=[]
            for para in shape.text_frame.paragraphs:
                runs=[]
                for run in para.runs:
                    f=run.font
                    runs.append({'text':run.text,'font':f.name,'size_emu':f.size,'bold':f.bold,'italic':f.italic,'underline':str(f.underline) if f.underline is not None else None})
                row['paragraphs'].append({'level':para.level,'alignment':str(para.alignment),'runs':runs})
        if shape.shape_type==MSO_SHAPE_TYPE.PICTURE:
            name=f'{trail}.{shape.image.ext}';(assets/name).write_bytes(shape.image.blob)
            row['image']={'path':'assets/'+name,'crop':{k:getattr(shape,'crop_'+k) for k in ['left','top','right','bottom']}}
        if shape.has_table:
            row['table']={'rows':[[cell.text for cell in r.cells] for r in shape.table.rows],
                          'column_widths_emu':[col.width for col in shape.table.columns]}
        if shape.has_chart:
            chart=shape.chart
            try:
                plots=[]
                for plot in chart.plots:
                    plots.append({'categories':[c.label for c in plot.categories],
                                  'series':[{'name':s.name,'values':list(s.values)} for s in plot.series]})
                row['chart']={'type':str(chart.chart_type),'plots':plots}
            except (AttributeError,ValueError,TypeError) as error:
                row['chart']={'type':str(chart.chart_type),'incomplete':str(error)}
                warnings.append(f'{trail}: chart properties incompletely extracted')
        recognized=any(k in row for k in ['children','text','image','table','chart'])
        if not recognized:
            row['unsupported']='Geometry retained; source XML/renderer needed for full appearance.'
            warnings.append(f'{trail}: appearance requires reconstruction from original source')
        return row
    slides=[]
    for n,slide in enumerate(prs.slides,1):
        slides.append({'number':n,'title':slide.shapes.title.text if slide.shapes.title else '',
                       'shapes':[shape_data(s,f'slide-{n}-shape-{i+1}') for i,s in enumerate(slide.shapes)],
                       'notes':slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ''})
    result={'source':str(Path(source).resolve()),'canvas_emu':{'width':prs.slide_width,'height':prs.slide_height},
            'slides':slides,'warnings':warnings,'scope':'Content/assets/geometry, not a faithful slide render or original-layer recovery'}
    (output/'extracted-slides.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    return {'file':str(output/'extracted-slides.json'),'slides':len(slides),'warnings':warnings}


if __name__=='__main__':
    if hasattr(sys.stdout,'reconfigure'):sys.stdout.reconfigure(encoding='utf-8')
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('pptx');parser.add_argument('output')
    options=parser.parse_args()
    try: print(json.dumps(extract(options.pptx,options.output),ensure_ascii=False))
    except ImportError: parser.exit(1,'python-pptx is required for PPTX extraction. Use an available document tool or install it in your chosen environment.\n')
