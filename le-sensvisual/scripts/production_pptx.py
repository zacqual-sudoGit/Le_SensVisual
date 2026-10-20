"""Native PPTX adapter for an authored production record; no Codex dependency.

Requires python-pptx 1.x (and its normal dependencies). Does not generate art,
render, infer layouts, resize text to fit, download packages or certify quality.
"""
import math
from pathlib import Path
import tempfile

from production import object_name, resolve


def colour(value, fallback='24362D'):
    text = str(value or fallback).lstrip('#')
    if len(text) != 6 or any(c not in '0123456789abcdefABCDEF' for c in text):
        raise ValueError('Use a six-digit RGB colour: ' + text)
    return text.upper()


def assemble(plan, plan_file, destination):
    try:
        from pptx import Presentation
        from pptx.chart.data import CategoryChartData
        from pptx.dml.color import RGBColor
        from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
        from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
        from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
        from pptx.oxml.xmlchemy import OxmlElement
        from pptx.util import Emu, Inches, Pt
        from PIL import Image
    except ImportError as error:
        raise ImportError('PPTX assembly needs python-pptx 1.x and its dependencies in this host. Use an available writer or the authorized package setup; no automatic installation was attempted.') from error

    pages = plan['pages']
    aspect = pages[0]['size'][0] / pages[0]['size'][1]
    if any(not math.isclose(p['size'][0] / p['size'][1], aspect, rel_tol=1e-6) for p in pages):
        raise ValueError('PPTX has one aspect ratio per deck. Resolve these output page proportions instead of stretching different canvases.')
    prs = Presentation()
    width_inches = plan.get('pptx_width_inches', 13.333333)
    if not isinstance(width_inches, (int, float)) or not math.isfinite(width_inches) or not 1 <= width_inches <= 56:
        raise ValueError('pptx_width_inches must be a finite supported presentation width (1–56)')
    prs.slide_width = Inches(width_inches)
    prs.slide_height = Emu(round(prs.slide_width / aspect))
    if not 914400 <= prs.slide_height <= 56 * 914400:
        raise ValueError('Requested PPTX proportions exceed supported physical slide dimensions')
    assets = {a['id']: a for a in plan.get('assets', [])}
    default_font = plan.get('font_family', 'Arial')

    for page in pages:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(colour(page.get('background'), 'FFFFFF'))
        k = prs.slide_width / page['size'][0]

        def length(v):
            return Emu(round(v * k))

        def font_size(v):
            return Pt(v * k / 12700)

        def set_alpha(element, opacity):
            if opacity < 1:
                alpha = OxmlElement('a:alpha')
                alpha.set('val', str(round(opacity * 100000)))
                element.append(alpha)

        def style(shape, obj):
            shape.shadow.inherit = False
            fill = obj.get('fill', 'none')
            if fill in {'none', None}:
                shape.fill.background()
            else:
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor.from_string(colour(fill))
                nodes = shape._element.xpath('.//a:solidFill/a:srgbClr')
                if nodes:
                    set_alpha(nodes[0], obj.get('opacity', 1))
            if obj.get('stroke') not in {None, 'none'}:
                shape.line.color.rgb = RGBColor.from_string(colour(obj['stroke']))
                shape.line.width = font_size(obj.get('stroke_width', 1))
                nodes = shape._element.xpath('.//a:ln/a:solidFill/a:srgbClr')
                if nodes:
                    set_alpha(nodes[0], obj.get('opacity', 1))
            else:
                shape.line.fill.background()

        def text(shape, value, obj):
            frame = shape.text_frame
            frame.clear()
            frame.word_wrap = True
            frame.auto_size = MSO_AUTO_SIZE.NONE
            # Table cells store insets/anchor on tcPr; text boxes use bodyPr.
            layout = shape if hasattr(shape, 'margin_left') else frame
            for side in ('left', 'right', 'top', 'bottom'):
                setattr(layout, 'margin_' + side, length(obj.get('padding', 0)))
            layout.vertical_anchor = {'top': MSO_ANCHOR.TOP, 'middle': MSO_ANCHOR.MIDDLE, 'bottom': MSO_ANCHOR.BOTTOM}.get(obj.get('vertical', 'top'), MSO_ANCHOR.TOP)
            for index, line in enumerate(str(value).split('\n')):
                p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                p.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}.get(obj.get('align', 'left'), PP_ALIGN.LEFT)
                p.space_before = Pt(0)
                p.space_after = Pt(0)
                p.line_spacing = float(obj.get('line_spacing', 1.15))
                run = p.add_run()
                run.text = line
                run.font.name = obj.get('font_family', default_font)
                run.font.size = font_size(obj.get('font_size', 24))
                run.font.bold = bool(obj.get('bold', False))
                run.font.italic = bool(obj.get('italic', False))
                run.font.color.rgb = RGBColor.from_string(colour(obj.get('color')))

        for obj in page['objects']:
            x, y, w, h = obj['box']
            name = object_name(page['id'], obj['id'])
            typ = obj['type']
            if typ == 'image':
                source = resolve(plan_file, assets[obj['asset']]['path'])
                with Image.open(source) as image:
                    iw, ih = image.size
                    image_format = image.format
                if image_format not in {'PNG', 'JPEG', 'GIF', 'BMP', 'TIFF'}:
                    raise ValueError('This PPTX adapter needs a library-supported raster asset. Preserve SVG/vector meaning through native paths or another capable host writer; do not silently drop ' + obj['id'])
                fit = obj.get('fit', 'contain')
                if fit == 'contain':
                    factor = min(w / iw, h / ih)
                    pw, ph = iw * factor, ih * factor
                    picture = slide.shapes.add_picture(str(source), length(x + (w - pw) / 2), length(y + (h - ph) / 2), length(pw), length(ph))
                else:
                    picture = slide.shapes.add_picture(str(source), length(x), length(y), length(w), length(h))
                    if fit == 'cover':
                        if iw / ih > w / h:
                            picture.crop_left = picture.crop_right = (1 - (w / h) / (iw / ih)) / 2
                        else:
                            picture.crop_top = picture.crop_bottom = (1 - (iw / ih) / (w / h)) / 2
                picture.name = name
                if obj.get('opacity', 1) < 1:
                    blip = picture._element.xpath('.//a:blip')[0]
                    alpha = OxmlElement('a:alphaModFix')
                    alpha.set('amt', str(round(obj['opacity'] * 100000)))
                    blip.append(alpha)
            elif typ == 'text':
                shape = slide.shapes.add_textbox(length(x), length(y), length(w), length(h))
                shape.name = name
                text(shape, obj['text'], obj)
            elif typ == 'shape':
                geometry = {'rect': MSO_SHAPE.RECTANGLE, 'round_rect': MSO_SHAPE.ROUNDED_RECTANGLE, 'ellipse': MSO_SHAPE.OVAL, 'triangle': MSO_SHAPE.ISOSCELES_TRIANGLE, 'chevron': MSO_SHAPE.CHEVRON}[obj.get('geometry', 'rect')]
                shape = slide.shapes.add_shape(geometry, length(x), length(y), length(w), length(h))
                shape.name = name
                if obj.get('geometry') == 'round_rect' and 'radius' in obj:
                    shape.adjustments[0] = min(0.5, max(0, obj['radius'] / min(w, h)))
                style(shape, obj)
            elif typ == 'polygon':
                points = obj['points']
                builder = slide.shapes.build_freeform(points[0][0], points[0][1], scale=k)
                builder.add_line_segments(points[1:], close=True)
                shape = builder.convert_to_shape()
                shape.name = name
                style(shape, obj)
            elif typ == 'line':
                points = obj['points']
                for index, (a, b) in enumerate(zip(points, points[1:])):
                    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, length(a[0]), length(a[1]), length(b[0]), length(b[1]))
                    shape.name = name if index == 0 else name + '-segment-' + str(index)
                    shape.line.color.rgb = RGBColor.from_string(colour(obj.get('stroke')))
                    shape.line.width = font_size(obj.get('stroke_width', 2))
                    if obj.get('arrow') and index == len(points) - 2:
                        end = OxmlElement('a:tailEnd')
                        end.set('type', 'triangle')
                        shape._element.spPr.get_or_add_ln().append(end)
            elif typ == 'table':
                rows = obj['rows']
                shape = slide.shapes.add_table(len(rows), len(rows[0]), length(x), length(y), length(w), length(h))
                shape.name = name
                for ri, row in enumerate(rows):
                    for ci, value in enumerate(row):
                        cell = shape.table.cell(ri, ci)
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = RGBColor.from_string(colour(obj.get('header_fill') if ri == 0 else obj.get('fill'), 'FFFFFF'))
                        text(cell, value, {**obj, 'font_size': obj.get('font_size', 20), 'bold': ri == 0, 'padding': obj.get('padding', 8)})
            elif typ == 'chart':
                data = CategoryChartData()
                data.categories = [str(c) for c in obj['categories']]
                for series in obj['series']:
                    data.add_series(str(series.get('name', '')), series['values'])
                kind = {'bar': XL_CHART_TYPE.BAR_CLUSTERED, 'column': XL_CHART_TYPE.COLUMN_CLUSTERED, 'line': XL_CHART_TYPE.LINE}[obj['kind']]
                shape = slide.shapes.add_chart(kind, length(x), length(y), length(w), length(h), data)
                shape.name = name
                chart = shape.chart
                chart.has_legend = len(obj['series']) > 1
                if chart.has_legend:
                    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
                    chart.legend.font.name = obj.get('font_family', default_font)
                    chart.legend.font.size = font_size(obj.get('font_size', 18))
                for axis in (chart.category_axis, chart.value_axis):
                    axis.tick_labels.font.name = obj.get('font_family', default_font)
                    axis.tick_labels.font.size = font_size(obj.get('font_size', 18))
                for series, source in zip(chart.series, obj['series']):
                    if source.get('color'):
                        if obj['kind'] == 'line':
                            series.format.line.color.rgb = RGBColor.from_string(colour(source['color']))
                        else:
                            series.format.fill.solid()
                            series.format.fill.fore_color.rgb = RGBColor.from_string(colour(source['color']))
        if page.get('notes'):
            slide.notes_slide.notes_text_frame.text = page['notes']

    # Write beside the destination and replace only after the writer completes.
    destination = Path(destination)
    with tempfile.NamedTemporaryFile(prefix='.visuasense-', suffix='.pptx', dir=destination.parent, delete=False) as f:
        temporary = Path(f.name)
    try:
        prs.save(str(temporary))
        temporary.replace(destination)
    finally:
        if temporary.exists():
            temporary.unlink()
